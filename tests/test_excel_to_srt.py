"""The Excel → SRT conversion, against a real Language Reactor export."""
from pathlib import Path

import pytest

from lib.subs import excel_to_srt
from lib.config import UserError, settings
from lib.subs.excel_to_srt import build_segments, ms_to_srt, parse_time

DATA = Path(__file__).parent / "data"


def test_parse_time_handles_every_shape_lr_emits():
    assert parse_time("42s") == 42000
    assert parse_time("2:14") == 134000
    assert parse_time("1:02:14.5") == 3734500
    assert parse_time(None) is None
    assert parse_time("") is None


def test_ms_to_srt():
    assert ms_to_srt(3734500) == "01:02:14,500"
    # LR's cue times are JS numbers: a fractional one must format, not raise.
    assert ms_to_srt(3734500.0000000005) == "01:02:14,500"


def test_segments_never_overlap_the_next_line():
    times = [0, 1000, 100000]
    segments = build_segments(times, ["a very long line of dialogue " * 5, "b", "c"])

    assert segments[0].end_ms < segments[1].start_ms  # trimmed to fit, however long the text
    assert segments[1].end_ms > segments[1].start_ms


def test_rows_with_no_timestamp_are_dropped():
    segments = build_segments([0, None, 5000], ["a", "b", "c"])
    assert segments[1] is None

    srt = excel_to_srt.create_srt_text(segments, ["a", "b", "c"])
    assert "b" not in srt
    assert srt.startswith("1\n00:00:00,000 --> ")


def test_real_export_produces_two_aligned_srt_files(tmp_path):
    primary, secondary = tmp_path / "p.srt", tmp_path / "s.srt"
    excel_to_srt.convert(DATA / "lln_excel_subs_sample.xlsx", primary, secondary)

    def timings(path):
        return [line for line in path.read_text(encoding="utf-8").splitlines() if " --> " in line]

    assert timings(primary), "no subtitles came out of the export"
    # The point of the pair: the two tracks share timing, so they line up in a player.
    assert timings(primary) == timings(secondary)


# ── verify: what counts as a complete export ──


def _xlsx_rows(path, rows, header=("Time", "Subtitle", "Machine Translation")):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.append(header)
    for row in rows:
        wb.active.append(row)
    wb.save(path)


def test_merged_rows_are_not_truncation(tmp_path):
    """LR merges consecutive cues into sentence rows — a 1220-cue track exports 862 rows — so
    an export with far fewer rows than cues is complete as long as it reaches the last cue.
    Counting rows instead rejected good exports and re-ripped episodes forever."""
    src = tmp_path / "e.xlsx"
    _xlsx_rows(src, [("0:01", "Hallo", "Hello"), ("1:04:16", "Tschüss", "Bye")])

    ok, why = excel_to_srt.verify(src, expected_end_ms=3_856_394)  # 1:04:16, the last cue

    assert ok, why


def test_an_export_that_stops_early_is_rejected(tmp_path):
    src = tmp_path / "e.xlsx"
    _xlsx_rows(src, [("0:01", "Hallo", "Hello"), ("10:00", "Tschüss", "Bye")])

    ok, why = excel_to_srt.verify(src, expected_end_ms=3_856_394)

    assert not ok
    assert "stops at 00:10:00,000" in why


def test_the_machine_translation_column_wins_over_the_human_one(tmp_path):
    """A title with official subtitles in your language exports 'Human Translation' first —
    but that column is empty for an ASR track, so matching it emptied the Secondary SRT."""
    src = tmp_path / "e.xlsx"
    _xlsx_rows(src, [("0:01", "Hallo", None, "Hello")],
               header=("Time", "Subtitle", "Human Translation", "Machine Translation"))

    ok, _ = excel_to_srt.verify(src)
    primary, secondary = tmp_path / "p.srt", tmp_path / "s.srt"
    excel_to_srt.convert(src, primary, secondary)

    assert ok, "the machine translation is there, so nothing is missing"
    assert "Hello" in secondary.read_text(encoding="utf-8")


# ── convert_exports: the xlsx is the only copy, so deleting it is the dangerous part ──


def _xlsx(path, rows):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.append(("Time", "Subtitle", "Machine Translation"))
    for row in rows:
        wb.active.append(row)
    wb.save(path)


@pytest.fixture
def downloads(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "EXPORTED_FILES_DIR", str(tmp_path))
    monkeypatch.setattr(settings, "RM_PROCESSED_FILES", True)
    # Pin the output/trash dirs too. Every Settings field reads GIGAKU_<NAME>, so a developer
    # with GIGAKU_SRT_TARGET_DIR exported (the normal setup) otherwise had these tests write
    # into their real subtitle library and fail on an empty tmp_path. Tests that care about
    # these dirs set them themselves, after this fixture.
    monkeypatch.setattr(settings, "SRT_TARGET_DIR", "")
    monkeypatch.setattr(settings, "SRT_TRASH_DIR", "")
    return tmp_path


def test_complete_exports_are_converted_and_only_then_deleted(downloads):
    src = downloads / "lln_excel_subs_1.xlsx"
    _xlsx(src, [("0:01", "Hallo", "Hello"), ("0:03", "Tschüss", "Bye")])

    excel_to_srt.convert_exports("Dark", season=2)

    assert (downloads / "Dark" / "Dark - S02E01 - Primary.srt").exists()
    assert (downloads / "Dark" / "Dark - S02E01 - Secondary.srt").exists()
    assert not src.exists()  # consumed, because it passed


def test_an_incomplete_export_is_still_converted_but_kept(downloads):
    """An export with missing translations still yields a perfectly good Primary track, so it
    is converted (polyglotka always did) — but it is *not* deleted, because the xlsx is the
    only copy of the lines its SRT is missing."""
    src = downloads / "lln_excel_subs_1.xlsx"
    _xlsx(src, [("0:01", "Hallo", "Hello"), ("0:03", "Tschüss", None)])

    written = excel_to_srt.convert_exports("Dark", season=2)

    primary = downloads / "Dark" / "Dark - S02E01 - Primary.srt"
    assert written == [str(primary)]
    assert primary.read_text(encoding="utf-8").count(" --> ") == 2  # both lines converted
    assert src.exists(), "an unverified export must survive"


def test_an_export_with_no_translation_column_still_converts(downloads):
    """LR can export subtitles without machine translation at all. polyglotka wrote the primary
    SRT anyway; refusing the file outright would be a regression."""
    import openpyxl

    src = downloads / "lln_excel_subs_1.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(("Time", "Subtitle"))
    wb.active.append(("0:01", "Hallo"))
    wb.save(src)

    written = excel_to_srt.convert_exports("Dark", season=2)

    assert len(written) == 1
    assert (downloads / "Dark" / "Dark - S02E01 - Primary.srt").exists()
    assert not (downloads / "Dark" / "Dark - S02E01 - Secondary.srt").exists()
    assert src.exists()  # nothing to verify against, so nothing gets deleted


def test_srt_target_dir_sends_the_output_to_the_subtitle_library(downloads, tmp_path, monkeypatch):
    library = tmp_path / "german_subs"
    monkeypatch.setattr(settings, "SRT_TARGET_DIR", str(library))
    _xlsx(downloads / "lln_excel_subs_1.xlsx", [("0:01", "Hallo", "Hello")])

    excel_to_srt.convert_exports("Dark", season=2)

    assert (library / "Dark" / "Dark - S02E01 - Primary.srt").exists()
    assert not (downloads / "Dark").exists()


def test_existing_srts_are_moved_to_the_trash_before_a_fresh_run(downloads, tmp_path, monkeypatch):
    """Re-ripping a shorter season must not leave last week's E09 sitting beside the new E01."""
    trash = tmp_path / "trash_subs"
    monkeypatch.setattr(settings, "SRT_TRASH_DIR", str(trash))
    season = downloads / "Dark"
    season.mkdir()
    stale = season / "Dark - S02E09 - Primary.srt"
    stale.write_text("old", encoding="utf-8")
    _xlsx(downloads / "lln_excel_subs_1.xlsx", [("0:01", "Hallo", "Hello")])

    excel_to_srt.convert_exports("Dark", season=2)

    assert not stale.exists()
    assert (trash / "Dark - S02E09 - Primary.srt").read_text(encoding="utf-8") == "old"
    assert (season / "Dark - S02E01 - Primary.srt").exists()  # the new one is written


def test_season_0_refuses_to_collapse_several_exports_onto_one_name(downloads):
    """Without an episode number every file resolves to the same basename: N exports would
    overwrite one SRT pair, and then all N originals would be deleted."""
    for i in (1, 2):
        _xlsx(downloads / f"lln_excel_subs_{i}.xlsx", [("0:01", "Hallo", "Hello")])

    with pytest.raises(UserError, match="overwrite each other"):
        excel_to_srt.convert_exports("Movie", season=0)

    assert len(list(downloads.glob("lln_excel_subs_*.xlsx"))) == 2  # nothing destroyed


def test_a_title_cannot_escape_the_downloads_directory(downloads):
    _xlsx(downloads / "lln_excel_subs_1.xlsx", [("0:01", "Hallo", "Hello")])

    excel_to_srt.convert_exports("../evil", season=0)

    assert not (downloads.parent / "evil").exists()
    assert list(downloads.glob("*/*.srt")), "the SRT stays under EXPORTED_FILES_DIR"


def test_verify_without_a_translation_column(tmp_path):
    """`gigaku subs` no longer asks LR for machine translation, so the export has no
    translation column at all — the end-timestamp check is what completeness rests on."""
    path = tmp_path / "no-mt.xlsx"
    _xlsx_rows(path, [("0s", "Eins"), ("10s", "Zwei")], header=("Time", "Subtitle"))
    assert excel_to_srt.verify(path, require_translation=False)[0]
    assert not excel_to_srt.verify(path)[0]  # the old contract still demands one


def test_an_export_of_the_wrong_track_is_rejected_by_its_own_text(tmp_path):
    """The E11/E12 failure: LR reverted to the title's default Japanese captions between the
    last track verify and the export click, so a file every other check passed was Japanese.
    Nothing upstream reads the text; this does, and it is what stops the xlsx being deleted
    and the SRT being written under a German name."""
    path = tmp_path / "jp.xlsx"
    _xlsx_rows(path, [(f"{i}s", "これは日本語の文章です。ひらがなとカタカナと漢字。") for i in range(12)],
               header=("Time", "Subtitle"))

    ok, why = excel_to_srt.verify(path, require_translation=False, language="German")

    assert not ok
    assert "wrong track" in why and "Japanese" in why


def test_the_right_track_passes_and_no_language_checks_nothing(tmp_path):
    path = tmp_path / "de.xlsx"
    _xlsx_rows(path, [(f"{i}s", "Die Akademie ist halb zerstört, also haben wir Ferien.")
                      for i in range(12)], header=("Time", "Subtitle"))
    assert excel_to_srt.verify(path, require_translation=False, language="German")[0]

    jp = tmp_path / "jp2.xlsx"
    _xlsx_rows(jp, [(f"{i}s", "これは日本語の文章です。ひらがなとカタカナと漢字。") for i in range(12)],
               header=("Time", "Subtitle"))
    # `gigaku srt` passes no language — it can't know one — and must behave exactly as before.
    assert excel_to_srt.verify(jp, require_translation=False)[0]


def test_convert_writes_the_primary_alone(tmp_path):
    path = tmp_path / "no-mt.xlsx"
    _xlsx_rows(path, [("0s", "Eins")], header=("Time", "Subtitle"))
    primary = tmp_path / "p.srt"
    excel_to_srt.convert(path, primary)
    assert "Eins" in primary.read_text(encoding="utf-8")
    assert not (tmp_path / "s.srt").exists()
